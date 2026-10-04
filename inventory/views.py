from decimal import Decimal, InvalidOperation
from io import BytesIO

from django.contrib.auth import authenticate, get_user_model
from django.db import transaction
from django.db.models import Q
from django.http import HttpResponse
from django.utils import timezone

from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from rest_framework import status
from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework.permissions import (
    AllowAny,
    IsAuthenticated,
)
from rest_framework.pagination import PageNumberPagination
from rest_framework.parsers import (
    MultiPartParser,
    FormParser,
    JSONParser,
)

from rest_framework_simplejwt.tokens import RefreshToken

from .models import (
    Branch,
    Product,
    Supplier,
    PurchaseOrder,
    PurchaseOrderItem,
    PurchaseOrderContainer,
    PurchaseReceipt,
    PurchaseReceiptItem,
)

from .serializers import (
    BranchSerializer,
    ProductSerializer,
    SupplierSerializer,
    PurchaseOrderListSerializer,
    PurchaseOrderDetailSerializer,
    PurchaseOrderWriteSerializer,
    PurchaseOrderContainerSerializer,
    PurchaseReceiptSerializer,
    BulkDeleteSerializer,
)


User = get_user_model()


# ============================================================
# USER / BRANCH HELPERS
# ============================================================

def get_user_type(user):

    if user.is_superuser:
        return "admin"

    profile = getattr(
        user,
        "profile",
        None
    )

    if profile is None:
        return None

    return profile.user_type


def get_assigned_branch(user):

    if user.is_superuser:
        return None

    profile = getattr(
        user,
        "profile",
        None
    )

    if profile is None:
        return None

    return profile.branch


def get_allowed_branches(user):

    user_type = get_user_type(user)

    if (
        user.is_superuser
        or user_type == "admin"
    ):
        return Branch.objects.all()

    if user_type == "staff":

        branch = get_assigned_branch(
            user
        )

        if branch:
            return Branch.objects.filter(
                pk=branch.pk
            )

    return Branch.objects.none()


def get_selected_branch(
    user,
    branch_id
):

    if not branch_id:
        return None

    try:
        return get_allowed_branches(
            user
        ).get(
            pk=branch_id
        )

    except (
        Branch.DoesNotExist,
        ValueError,
        TypeError,
    ):
        return None


def user_payload(user):

    user_type = get_user_type(
        user
    )

    branch = get_assigned_branch(
        user
    )

    return {
        "id": user.id,

        "username":
        user.username,

        "name":
        user.get_full_name()
        or user.username,

        "user_type":
        user_type,

        "branch":
        (
            {
                "id":
                branch.id,

                "branch_name":
                branch.branch_name,
            }
            if branch
            else None
        ),

        "requires_branch_selection":
        (
            user_type == "admin"
            and not user.is_superuser
        ),
    }


def get_allowed_products(user):

    user_type = get_user_type(
        user
    )

    queryset = (
        Product.objects
        .select_related("branch")
        .all()
    )

    if (
        user.is_superuser
        or user_type == "admin"
    ):
        return queryset

    if user_type == "staff":

        branch = get_assigned_branch(
            user
        )

        if branch:
            return queryset.filter(
                branch=branch
            )

    return Product.objects.none()


def resolve_product_branch(
    user,
    branch_id=None
):

    user_type = get_user_type(
        user
    )

    if (
        user.is_superuser
        or user_type == "admin"
    ):

        if not branch_id:
            return None

        return get_selected_branch(
            user,
            branch_id
        )

    if user_type == "staff":
        return get_assigned_branch(
            user
        )

    return None


# ============================================================
# LOGIN
# ============================================================

class LoginView(APIView):

    permission_classes = [
        AllowAny
    ]

    authentication_classes = []

    def post(
        self,
        request
    ):

        username = request.data.get(
            "username",
            ""
        ).strip()

        password = request.data.get(
            "password",
            ""
        )

        if not username or not password:

            return Response(
                {
                    "detail":
                    "Username and password are required."
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        user = authenticate(
            request,
            username=username,
            password=password,
        )

        if user is None:

            return Response(
                {
                    "detail":
                    "Invalid username or password."
                },
                status=status.HTTP_401_UNAUTHORIZED,
            )

        if not user.is_active:

            return Response(
                {
                    "detail":
                    "This account is inactive."
                },
                status=status.HTTP_403_FORBIDDEN,
            )

        if (
            not user.is_superuser
            and not hasattr(
                user,
                "profile"
            )
        ):

            return Response(
                {
                    "detail":
                    "User profile is not configured."
                },
                status=status.HTTP_403_FORBIDDEN,
            )

        if get_user_type(user) not in (
            "admin",
            "staff",
        ):

            return Response(
                {
                    "detail":
                    "This user has no valid user type."
                },
                status=status.HTTP_403_FORBIDDEN,
            )

        refresh = RefreshToken.for_user(
            user
        )

        return Response({
            "access":
            str(refresh.access_token),

            "refresh":
            str(refresh),

            "user":
            user_payload(user),

            "branches":
            BranchSerializer(
                get_allowed_branches(user),
                many=True,
            ).data,
        })


# ============================================================
# LOGOUT
# ============================================================

class LogoutView(APIView):

    permission_classes = [
        IsAuthenticated
    ]

    def post(
        self,
        request
    ):

        refresh_token = request.data.get(
            "refresh"
        )

        if not refresh_token:

            return Response(
                {
                    "detail":
                    "Refresh token is required."
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        try:

            token = RefreshToken(
                refresh_token
            )

            token.blacklist()

        except Exception:

            return Response(
                {
                    "detail":
                    "Invalid refresh token."
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        return Response({
            "detail":
            "Logged out successfully."
        })


# ============================================================
# ME
# ============================================================

class MeView(APIView):

    permission_classes = [
        IsAuthenticated
    ]

    def get(
        self,
        request
    ):

        return Response({
            "user":
            user_payload(
                request.user
            )
        })


# ============================================================
# BRANCH LIST
# ============================================================

class BranchListView(APIView):

    permission_classes = [
        IsAuthenticated
    ]

    def get(
        self,
        request
    ):

        branches = get_allowed_branches(
            request.user
        )

        return Response(
            BranchSerializer(
                branches,
                many=True
            ).data
        )


# ============================================================
# BRANCH SELECT
# ============================================================

class BranchSelectView(APIView):

    permission_classes = [
        IsAuthenticated
    ]

    def post(
        self,
        request
    ):

        user = request.user

        user_type = get_user_type(
            user
        )

        if user_type not in (
            "admin",
            "staff",
        ):

            return Response(
                {
                    "detail":
                    "Access denied."
                },
                status=status.HTTP_403_FORBIDDEN,
            )

        branch_id = request.data.get(
            "branch_id"
        )

        branch = get_selected_branch(
            user,
            branch_id
        )

        if branch is None:

            return Response(
                {
                    "detail":
                    "Invalid or unauthorized branch."
                },
                status=status.HTTP_403_FORBIDDEN,
            )

        return Response({
            "detail":
            "Branch selected successfully.",

            "branch":
            BranchSerializer(
                branch
            ).data,
        })


# ============================================================
# BRANCH DETAILS
# ============================================================

class BranchDetailsView(APIView):

    permission_classes = [
        IsAuthenticated
    ]

    def get(
        self,
        request
    ):

        branch_id = request.query_params.get(
            "branch_id"
        )

        if not branch_id:

            return Response(
                {
                    "detail":
                    "branch_id is required."
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        branch = get_selected_branch(
            request.user,
            branch_id
        )

        if branch is None:

            return Response(
                {
                    "detail":
                    "Branch not found or access denied."
                },
                status=status.HTTP_403_FORBIDDEN,
            )

        return Response(
            BranchSerializer(
                branch
            ).data
        )


# ============================================================
# PAGINATION
# ============================================================

class ProductPagination(
    PageNumberPagination
):

    page_size = 50

    page_size_query_param = (
        "page_size"
    )

    max_page_size = 500


class PurchaseOrderPagination(
    PageNumberPagination
):

    page_size = 20

    page_size_query_param = (
        "page_size"
    )

    max_page_size = 100


# ============================================================
# PRODUCT LIST + CREATE
# ============================================================

class ProductListCreateView(APIView):

    permission_classes = [
        IsAuthenticated
    ]

    parser_classes = [
        JSONParser,
        MultiPartParser,
        FormParser,
    ]

    def get(
        self,
        request
    ):

        queryset = get_allowed_products(
            request.user
        )

        search = request.query_params.get(
            "search",
            ""
        ).strip()

        barcode = request.query_params.get(
            "barcode",
            ""
        ).strip()

        short_description = (
            request.query_params.get(
                "short_description",
                ""
            ).strip()
        )

        location = request.query_params.get(
            "location",
            ""
        ).strip()

        branch_id = request.query_params.get(
            "branch_id"
        )

        if search:

            queryset = queryset.filter(
                Q(
                    barcode__icontains=search
                )
                |
                Q(
                    short_description__icontains=search
                )
            )

        if barcode:

            queryset = queryset.filter(
                barcode__icontains=barcode
            )

        if short_description:

            queryset = queryset.filter(
                short_description__icontains=(
                    short_description
                )
            )

        if location:

            queryset = queryset.filter(
                location__icontains=location
            )

        if branch_id:

            branch = get_selected_branch(
                request.user,
                branch_id
            )

            if branch is None:

                return Response(
                    {
                        "detail":
                        "Invalid or unauthorized branch."
                    },
                    status=status.HTTP_403_FORBIDDEN,
                )

            queryset = queryset.filter(
                branch=branch
            )

        ordering = request.query_params.get(
            "ordering",
            "barcode"
        )

        allowed_ordering = {
            "barcode",
            "-barcode",
            "location",
            "-location",
            "cost_price",
            "-cost_price",
            "selling_price",
            "-selling_price",
            "stk_on_hand",
            "-stk_on_hand",
            "physical_qty",
            "-physical_qty",
            "created_at",
            "-created_at",
            "updated_at",
            "-updated_at",
        }

        if ordering not in allowed_ordering:
            ordering = "barcode"

        queryset = queryset.order_by(
            ordering
        )

        paginator = ProductPagination()

        page = paginator.paginate_queryset(
            queryset,
            request
        )

        serializer = ProductSerializer(
            page,
            many=True
        )

        return paginator.get_paginated_response(
            serializer.data
        )

    def post(
        self,
        request
    ):

        user = request.user

        branch_id = request.data.get(
            "branch"
        )

        branch = resolve_product_branch(
            user,
            branch_id
        )

        if branch is None:

            return Response(
                {
                    "detail":
                    "A valid branch is required."
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        data = request.data.copy()

        data["branch"] = branch.id

        serializer = ProductSerializer(
            data=data
        )

        if not serializer.is_valid():

            return Response(
                serializer.errors,
                status=status.HTTP_400_BAD_REQUEST,
            )

        barcode = serializer.validated_data[
            "barcode"
        ]

        if Product.objects.filter(
            branch=branch,
            barcode=barcode
        ).exists():

            return Response(
                {
                    "detail":
                    "A product with this barcode "
                    "already exists in this branch."
                },
                status=status.HTTP_409_CONFLICT,
            )

        product = serializer.save(
            branch=branch
        )

        return Response(
            ProductSerializer(
                product
            ).data,
            status=status.HTTP_201_CREATED,
        )


# ============================================================
# PRODUCT DETAIL
# ============================================================

class ProductDetailView(APIView):

    permission_classes = [
        IsAuthenticated
    ]

    def get_object(
        self,
        request,
        pk
    ):

        try:

            return get_allowed_products(
                request.user
            ).get(
                pk=pk
            )

        except Product.DoesNotExist:

            return None

    def get(
        self,
        request,
        pk
    ):

        product = self.get_object(
            request,
            pk
        )

        if product is None:

            return Response(
                {
                    "detail":
                    "Product not found."
                },
                status=status.HTTP_404_NOT_FOUND,
            )

        return Response(
            ProductSerializer(
                product
            ).data
        )

    def put(
        self,
        request,
        pk
    ):

        return self._update(
            request,
            pk,
            partial=False
        )

    def patch(
        self,
        request,
        pk
    ):

        return self._update(
            request,
            pk,
            partial=True
        )

    def _update(
        self,
        request,
        pk,
        partial=False
    ):

        product = self.get_object(
            request,
            pk
        )

        if product is None:

            return Response(
                {
                    "detail":
                    "Product not found."
                },
                status=status.HTTP_404_NOT_FOUND,
            )

        data = request.data.copy()

        if (
            "branch" in data
            and str(data["branch"])
            != str(product.branch_id)
        ):

            return Response(
                {
                    "detail":
                    "Product branch cannot be changed."
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        data["branch"] = product.branch_id

        serializer = ProductSerializer(
            product,
            data=data,
            partial=partial
        )

        if not serializer.is_valid():

            return Response(
                serializer.errors,
                status=status.HTTP_400_BAD_REQUEST,
            )

        barcode = serializer.validated_data.get(
            "barcode",
            product.barcode
        )

        duplicate = Product.objects.filter(
            branch=product.branch,
            barcode=barcode
        ).exclude(
            pk=product.pk
        ).exists()

        if duplicate:

            return Response(
                {
                    "detail":
                    "Another product with this "
                    "barcode already exists."
                },
                status=status.HTTP_409_CONFLICT,
            )

        product = serializer.save()

        return Response(
            ProductSerializer(
                product
            ).data
        )

    def delete(
        self,
        request,
        pk
    ):

        product = self.get_object(
            request,
            pk
        )

        if product is None:

            return Response(
                {
                    "detail":
                    "Product not found."
                },
                status=status.HTTP_404_NOT_FOUND,
            )

        product.delete()

        return Response(
            {
                "detail":
                "Product deleted successfully."
            },
            status=status.HTTP_204_NO_CONTENT,
        )


# ============================================================
# PRODUCT BULK DELETE
# ============================================================

class ProductBulkDeleteView(APIView):

    permission_classes = [
        IsAuthenticated
    ]

    def post(
        self,
        request
    ):

        serializer = BulkDeleteSerializer(
            data=request.data
        )

        if not serializer.is_valid():

            return Response(
                serializer.errors,
                status=status.HTTP_400_BAD_REQUEST,
            )

        ids = serializer.validated_data[
            "ids"
        ]

        queryset = get_allowed_products(
            request.user
        ).filter(
            id__in=ids
        )

        deleted_count = queryset.count()

        queryset.delete()

        return Response({
            "detail":
            f"{deleted_count} product(s) deleted.",

            "deleted_count":
            deleted_count,
        })


# ============================================================
# EXCEL HELPERS
# ============================================================

EXPECTED_HEADERS = [
    "BarCode",
    "Short Description",
    "Location",
    "Cost Price",
    "Selling Price",
    "Stk On Hand",
    "Physical Qty",
]


def clean_excel_value(value):

    if value is None:
        return ""

    return str(value).strip()


def decimal_value(value):

    if value is None or value == "":
        return Decimal("0")

    try:

        return Decimal(
            str(value)
            .replace(",", "")
            .strip()
        )

    except (
        InvalidOperation,
        ValueError,
        TypeError,
    ):

        raise ValueError(
            f"Invalid decimal value: {value}"
        )


def integer_value(value):

    if value is None or value == "":
        return 0

    try:

        return int(
            float(
                str(value)
                .replace(",", "")
                .strip()
            )
        )

    except (
        ValueError,
        TypeError,
    ):

        raise ValueError(
            f"Invalid integer value: {value}"
        )


def normalize_headers(headers):

    return {
        str(value).strip(): index
        for index, value in enumerate(
            headers
        )
        if value is not None
    }


# ============================================================
# PRODUCT BULK IMPORT
# ============================================================

class ProductBulkImportView(APIView):

    permission_classes = [
        IsAuthenticated
    ]

    parser_classes = [
        MultiPartParser,
        FormParser,
    ]

    def post(
        self,
        request
    ):

        uploaded_file = request.FILES.get(
            "file"
        )

        if not uploaded_file:

            return Response(
                {
                    "detail":
                    "Excel file is required."
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        filename = uploaded_file.name.lower()

        if not filename.endswith(
            ".xlsx"
        ):

            return Response(
                {
                    "detail":
                    "Only .xlsx Excel files are supported."
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        branch_id = request.data.get(
            "branch_id"
        )

        branch = resolve_product_branch(
            request.user,
            branch_id
        )

        if branch is None:

            return Response(
                {
                    "detail":
                    "A valid branch is required."
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        try:

            workbook = load_workbook(
                uploaded_file,
                read_only=True,
                data_only=True,
            )

            if "Inventory" in workbook.sheetnames:

                worksheet = workbook[
                    "Inventory"
                ]

            else:

                worksheet = workbook[
                    workbook.sheetnames[0]
                ]

            rows = worksheet.iter_rows(
                values_only=True
            )

            try:

                headers = next(rows)

            except StopIteration:

                return Response(
                    {
                        "detail":
                        "Excel file is empty."
                    },
                    status=status.HTTP_400_BAD_REQUEST,
                )

            header_map = normalize_headers(
                headers
            )

            missing_headers = [
                header
                for header in EXPECTED_HEADERS
                if header not in header_map
            ]

            if missing_headers:

                return Response(
                    {
                        "detail":
                        "Required Excel columns are missing.",

                        "missing_columns":
                        missing_headers,

                        "expected_columns":
                        EXPECTED_HEADERS,
                    },
                    status=status.HTTP_400_BAD_REQUEST,
                )

            products_to_create = []

            products_to_update = []

            errors = []

            existing_products = {
                product.barcode: product
                for product in Product.objects.filter(
                    branch=branch
                )
            }

            file_barcodes = set()

            for row_number, row in enumerate(
                rows,
                start=2
            ):

                try:

                    barcode = clean_excel_value(
                        row[
                            header_map["BarCode"]
                        ]
                    )

                    if not barcode:

                        raise ValueError(
                            "Barcode is empty."
                        )

                    if barcode in file_barcodes:

                        raise ValueError(
                            "Duplicate barcode in Excel file."
                        )

                    file_barcodes.add(
                        barcode
                    )

                    short_description = (
                        clean_excel_value(
                            row[
                                header_map[
                                    "Short Description"
                                ]
                            ]
                        )
                    )

                    location = clean_excel_value(
                        row[
                            header_map[
                                "Location"
                            ]
                        ]
                    )

                    cost_price = decimal_value(
                        row[
                            header_map[
                                "Cost Price"
                            ]
                        ]
                    )

                    selling_price = decimal_value(
                        row[
                            header_map[
                                "Selling Price"
                            ]
                        ]
                    )

                    stk_on_hand = integer_value(
                        row[
                            header_map[
                                "Stk On Hand"
                            ]
                        ]
                    )

                    physical_qty = integer_value(
                        row[
                            header_map[
                                "Physical Qty"
                            ]
                        ]
                    )

                    if cost_price < 0:

                        raise ValueError(
                            "Cost Price cannot be negative."
                        )

                    if selling_price < 0:

                        raise ValueError(
                            "Selling Price cannot be negative."
                        )

                    if stk_on_hand < 0:

                        raise ValueError(
                            "Stk On Hand cannot be negative."
                        )

                    if physical_qty < 0:

                        raise ValueError(
                            "Physical Qty cannot be negative."
                        )

                    existing = (
                        existing_products.get(
                            barcode
                        )
                    )

                    if existing:

                        existing.short_description = (
                            short_description
                        )

                        existing.location = (
                            location
                        )

                        existing.cost_price = (
                            cost_price
                        )

                        existing.selling_price = (
                            selling_price
                        )

                        existing.stk_on_hand = (
                            stk_on_hand
                        )

                        existing.physical_qty = (
                            physical_qty
                        )

                        products_to_update.append(
                            existing
                        )

                    else:

                        products_to_create.append(
                            Product(
                                branch=branch,
                                barcode=barcode,
                                short_description=(
                                    short_description
                                ),
                                location=location,
                                cost_price=(
                                    cost_price
                                ),
                                selling_price=(
                                    selling_price
                                ),
                                stk_on_hand=(
                                    stk_on_hand
                                ),
                                physical_qty=(
                                    physical_qty
                                ),
                            )
                        )

                except Exception as exc:

                    errors.append({
                        "row":
                        row_number,

                        "error":
                        str(exc),
                    })

            if errors:

                return Response(
                    {
                        "detail":
                        "Import failed because some rows contain errors.",

                        "errors":
                        errors[:100],

                        "total_errors":
                        len(errors),
                    },
                    status=status.HTTP_400_BAD_REQUEST,
                )

            with transaction.atomic():

                if products_to_create:

                    Product.objects.bulk_create(
                        products_to_create,
                        batch_size=1000,
                    )

                if products_to_update:

                    Product.objects.bulk_update(
                        products_to_update,
                        fields=[
                            "short_description",
                            "location",
                            "cost_price",
                            "selling_price",
                            "stk_on_hand",
                            "physical_qty",
                        ],
                        batch_size=1000,
                    )

            return Response({
                "detail":
                "Excel import completed successfully.",

                "branch":
                BranchSerializer(
                    branch
                ).data,

                "created":
                len(products_to_create),

                "updated":
                len(products_to_update),

                "total":
                (
                    len(products_to_create)
                    +
                    len(products_to_update)
                ),
            })

        except Exception as exc:

            return Response(
                {
                    "detail":
                    f"Unable to process Excel file: {exc}"
                },
                status=status.HTTP_400_BAD_REQUEST,
            )


# ============================================================
# PRODUCT EXPORT
# ============================================================

class ProductExportView(APIView):

    permission_classes = [
        IsAuthenticated
    ]

    def get(
        self,
        request
    ):

        queryset = get_allowed_products(
            request.user
        )

        branch_id = request.query_params.get(
            "branch_id"
        )

        if branch_id:

            branch = get_selected_branch(
                request.user,
                branch_id
            )

            if branch is None:

                return Response(
                    {
                        "detail":
                        "Invalid or unauthorized branch."
                    },
                    status=status.HTTP_403_FORBIDDEN,
                )

            queryset = queryset.filter(
                branch=branch
            )

        search = request.query_params.get(
            "search",
            ""
        ).strip()

        if search:

            queryset = queryset.filter(
                Q(
                    barcode__icontains=search
                )
                |
                Q(
                    short_description__icontains=search
                )
            )

        barcode = request.query_params.get(
            "barcode",
            ""
        ).strip()

        if barcode:

            queryset = queryset.filter(
                barcode__icontains=barcode
            )

        short_description = (
            request.query_params.get(
                "short_description",
                ""
            ).strip()
        )

        if short_description:

            queryset = queryset.filter(
                short_description__icontains=(
                    short_description
                )
            )

        location = request.query_params.get(
            "location",
            ""
        ).strip()

        if location:

            queryset = queryset.filter(
                location__icontains=location
            )

        queryset = queryset.order_by(
            "barcode"
        )

        workbook = Workbook()

        worksheet = workbook.active

        worksheet.title = "Inventory"

        headers = [
            "S.R",
            "BarCode",
            "Short Description",
            "Location",
            "Cost Price",
            "Selling Price",
            "Stk On Hand",
            "Physical Qty",
            "Variance",
        ]

        worksheet.append(headers)

        header_fill = PatternFill(
            fill_type="solid",
            fgColor="D9EAF7",
        )

        for cell in worksheet[1]:

            cell.font = Font(
                bold=True
            )

            cell.fill = header_fill

            cell.alignment = Alignment(
                horizontal="center"
            )

        for index, product in enumerate(
            queryset.iterator(
                chunk_size=1000
            ),
            start=1
        ):

            worksheet.append([
                index,
                product.barcode,
                product.short_description,
                product.location,
                float(product.cost_price),
                float(product.selling_price),
                product.stk_on_hand,
                product.physical_qty,
                product.variance,
            ])

        widths = {
            1: 10,
            2: 25,
            3: 45,
            4: 20,
            5: 15,
            6: 15,
            7: 15,
            8: 15,
            9: 15,
        }

        for column, width in widths.items():

            worksheet.column_dimensions[
                get_column_letter(column)
            ].width = width

        output = BytesIO()

        workbook.save(output)

        output.seek(0)

        response = HttpResponse(
            output.getvalue(),
            content_type=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
        )

        response[
            "Content-Disposition"
        ] = (
            'attachment; '
            'filename="inventory_export.xlsx"'
        )

        return response


# ============================================================
# SUPPLIER MANAGEMENT
# ============================================================

class SupplierListCreateView(APIView):

    permission_classes = [
        IsAuthenticated
    ]

    def get(
        self,
        request
    ):

        queryset = Supplier.objects.all()

        search = request.query_params.get(
            "search",
            ""
        ).strip()

        supplier_name = (
            request.query_params.get(
                "supplier_name",
                ""
            ).strip()
        )

        location = (
            request.query_params.get(
                "location",
                ""
            ).strip()
        )

        supplier_status = (
            request.query_params.get(
                "status",
                ""
            ).strip()
        )

        if search:

            queryset = queryset.filter(
                Q(
                    supplier_name__icontains=search
                )
                |
                Q(
                    supplier_location__icontains=search
                )
            )

        if supplier_name:

            queryset = queryset.filter(
                supplier_name__icontains=(
                    supplier_name
                )
            )

        if location:

            queryset = queryset.filter(
                supplier_location__icontains=(
                    location
                )
            )

        if supplier_status in (
            "active",
            "inactive",
        ):

            queryset = queryset.filter(
                status=supplier_status
            )

        return Response(
            SupplierSerializer(
                queryset,
                many=True
            ).data
        )

    def post(
        self,
        request
    ):

        serializer = SupplierSerializer(
            data=request.data
        )

        if not serializer.is_valid():

            return Response(
                serializer.errors,
                status=status.HTTP_400_BAD_REQUEST,
            )

        supplier = serializer.save()

        return Response(
            SupplierSerializer(
                supplier
            ).data,
            status=status.HTTP_201_CREATED,
        )


class SupplierDetailView(APIView):

    permission_classes = [
        IsAuthenticated
    ]

    def get_object(
        self,
        pk
    ):

        try:

            return Supplier.objects.get(
                pk=pk
            )

        except Supplier.DoesNotExist:

            return None

    def get(
        self,
        request,
        pk
    ):

        supplier = self.get_object(
            pk
        )

        if supplier is None:

            return Response(
                {
                    "detail":
                    "Supplier not found."
                },
                status=status.HTTP_404_NOT_FOUND,
            )

        return Response(
            SupplierSerializer(
                supplier
            ).data
        )

    def put(
        self,
        request,
        pk
    ):

        return self.update(
            request,
            pk,
            partial=False
        )

    def patch(
        self,
        request,
        pk
    ):

        return self.update(
            request,
            pk,
            partial=True
        )

    def update(
        self,
        request,
        pk,
        partial=False
    ):

        supplier = self.get_object(
            pk
        )

        if supplier is None:

            return Response(
                {
                    "detail":
                    "Supplier not found."
                },
                status=status.HTTP_404_NOT_FOUND,
            )

        serializer = SupplierSerializer(
            supplier,
            data=request.data,
            partial=partial
        )

        if not serializer.is_valid():

            return Response(
                serializer.errors,
                status=status.HTTP_400_BAD_REQUEST,
            )

        supplier = serializer.save()

        return Response(
            SupplierSerializer(
                supplier
            ).data
        )

    def delete(
        self,
        request,
        pk
    ):

        supplier = self.get_object(
            pk
        )

        if supplier is None:

            return Response(
                {
                    "detail":
                    "Supplier not found."
                },
                status=status.HTTP_404_NOT_FOUND,
            )

        has_orders = PurchaseOrder.objects.filter(
            supplier=supplier
        ).exists()

        if has_orders:

            supplier.status = "inactive"

            supplier.save(
                update_fields=[
                    "status",
                    "updated_at",
                ]
            )

            return Response({
                "detail":
                "Supplier has Purchase Orders, so it was deactivated instead of deleted.",

                "status":
                "inactive",
            })

        supplier.delete()

        return Response(
            {
                "detail":
                "Supplier deleted successfully."
            },
            status=status.HTTP_204_NO_CONTENT,
        )


# ============================================================
# SUPPLIER ACTIVATE / DEACTIVATE
# ============================================================

class SupplierStatusView(APIView):

    permission_classes = [
        IsAuthenticated
    ]

    def patch(
        self,
        request,
        pk
    ):

        try:

            supplier = Supplier.objects.get(
                pk=pk
            )

        except Supplier.DoesNotExist:

            return Response(
                {
                    "detail":
                    "Supplier not found."
                },
                status=status.HTTP_404_NOT_FOUND,
            )

        supplier_status = request.data.get(
            "status"
        )

        if supplier_status not in (
            "active",
            "inactive",
        ):

            return Response(
                {
                    "detail":
                    "Status must be active or inactive."
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        supplier.status = supplier_status

        supplier.save(
            update_fields=[
                "status",
                "updated_at",
            ]
        )

        return Response(
            SupplierSerializer(
                supplier
            ).data
        )


# ============================================================
# PO NUMBER GENERATOR
# ============================================================

def generate_po_number():

    current_year = timezone.now().year

    prefix = f"PO-{current_year}-"

    last_po = (
        PurchaseOrder.objects
        .filter(
            po_number__startswith=prefix
        )
        .order_by("-id")
        .first()
    )

    if not last_po:

        next_number = 1

    else:

        try:

            last_number = int(
                last_po.po_number.split(
                    "-"
                )[-1]
            )

            next_number = (
                last_number + 1
            )

        except (
            ValueError,
            IndexError,
        ):

            next_number = (
                PurchaseOrder.objects.filter(
                    po_number__startswith=prefix
                ).count()
                + 1
            )

    return (
        f"{prefix}"
        f"{next_number:05d}"
    )


# ============================================================
# PURCHASE ORDER PRODUCT SEARCH
# ============================================================

class PurchaseOrderProductSearchView(
    APIView
):

    permission_classes = [
        IsAuthenticated
    ]

    def get(
        self,
        request
    ):

        queryset = get_allowed_products(
            request.user
        )

        branch_id = request.query_params.get(
            "branch_id"
        )

        if branch_id:

            branch = get_selected_branch(
                request.user,
                branch_id
            )

            if branch is None:

                return Response(
                    {
                        "detail":
                        "Invalid or unauthorized branch."
                    },
                    status=status.HTTP_403_FORBIDDEN,
                )

            queryset = queryset.filter(
                branch=branch
            )

        barcode = request.query_params.get(
            "barcode",
            ""
        ).strip()

        short_description = (
            request.query_params.get(
                "short_description",
                ""
            ).strip()
        )

        location = request.query_params.get(
            "location",
            ""
        ).strip()

        if barcode:

            queryset = queryset.filter(
                barcode__icontains=barcode
            )

        if short_description:

            queryset = queryset.filter(
                short_description__icontains=(
                    short_description
                )
            )

        if location:

            queryset = queryset.filter(
                location__icontains=location
            )

        queryset = queryset.order_by(
            "barcode"
        )

        return Response(
            ProductSerializer(
                queryset[:100],
                many=True
            ).data
        )


# ============================================================
# PURCHASE ORDER QUERYSET
# ============================================================

def get_allowed_purchase_orders(
    user
):

    queryset = (
        PurchaseOrder.objects
        .select_related(
            "branch",
            "supplier",
            "created_by",
        )
        .prefetch_related(
            "items__product",
            "containers",
        )
    )

    user_type = get_user_type(
        user
    )

    if (
        user.is_superuser
        or user_type == "admin"
    ):

        return queryset

    if user_type == "staff":

        branch = get_assigned_branch(
            user
        )

        if branch:

            return queryset.filter(
                branch=branch
            )

    return queryset.none()


# ============================================================
# PURCHASE ORDER LIST
# ============================================================

class PurchaseOrderListView(APIView):

    permission_classes = [
        IsAuthenticated
    ]

    def get(
        self,
        request
    ):

        queryset = get_allowed_purchase_orders(
            request.user
        )

        branch_id = request.query_params.get(
            "branch_id"
        )

        if branch_id:

            branch = get_selected_branch(
                request.user,
                branch_id
            )

            if branch is None:

                return Response(
                    {
                        "detail":
                        "Invalid or unauthorized branch."
                    },
                    status=status.HTTP_403_FORBIDDEN,
                )

            queryset = queryset.filter(
                branch=branch
            )

        supplier_id = request.query_params.get(
            "supplier_id"
        )

        if supplier_id:

            queryset = queryset.filter(
                supplier_id=supplier_id
            )

        po_status = request.query_params.get(
            "status"
        )

        if po_status:

            queryset = queryset.filter(
                status=po_status
            )

        search = request.query_params.get(
            "search",
            ""
        ).strip()

        if search:

            queryset = queryset.filter(
                Q(
                    po_number__icontains=search
                )
                |
                Q(
                    supplier__supplier_name__icontains=search
                )
            )

        queryset = queryset.order_by(
            "-created_at"
        )

        paginator = PurchaseOrderPagination()

        page = paginator.paginate_queryset(
            queryset,
            request
        )

        serializer = (
            PurchaseOrderListSerializer(
                page,
                many=True
            )
        )

        return paginator.get_paginated_response(
            serializer.data
        )


# ============================================================
# PURCHASE ORDER CREATE
# ============================================================

class PurchaseOrderCreateView(APIView):

    permission_classes = [
        IsAuthenticated
    ]

    def post(
        self,
        request
    ):

        user = request.user

        serializer = PurchaseOrderWriteSerializer(
            data=request.data
        )

        if not serializer.is_valid():

            return Response(
                serializer.errors,
                status=status.HTTP_400_BAD_REQUEST,
            )

        validated = serializer.validated_data

        branch_id = validated.get(
            "branch_id"
        )

        branch = resolve_product_branch(
            user,
            branch_id
        )

        if branch is None:

            return Response(
                {
                    "detail":
                    "A valid branch is required."
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        supplier = validated[
            "supplier"
        ]

        if supplier.status != "active":

            return Response(
                {
                    "detail":
                    "Inactive suppliers cannot be used for new Purchase Orders."
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        items = validated[
            "items"
        ]

        containers = validated.get(
            "containers",
            []
        )

        for item in items:

            product = item[
                "product"
            ]

            if product.branch_id != branch.id:

                return Response(
                    {
                        "detail":
                        (
                            f"Product "
                            f"{product.barcode} "
                            f"does not belong to selected branch."
                        )
                    },
                    status=status.HTTP_400_BAD_REQUEST,
                )

        with transaction.atomic():

            po_number = generate_po_number()

            purchase_order = PurchaseOrder.objects.create(
                po_number=po_number,
                branch=branch,
                supplier=supplier,
                po_date=validated[
                    "po_date"
                ],
                expected_delivery_date=(
                    validated.get(
                        "expected_delivery_date"
                    )
                ),
                status="draft",
                created_by=user,
            )

            for item in items:

                PurchaseOrderItem.objects.create(
                    purchase_order=purchase_order,
                    product=item["product"],
                    quantity=item["quantity"],
                    purchase_price=item[
                        "purchase_price"
                    ],
                )

            for container in containers:

                PurchaseOrderContainer.objects.create(
                    purchase_order=purchase_order,
                    container_number=container[
                        "container_number"
                    ],
                    container_type=container.get(
                        "container_type",
                        ""
                    ),
                    quantity=container[
                        "quantity"
                    ],
                    remarks=container.get(
                        "remarks",
                        ""
                    ),
                )

        purchase_order = (
            get_allowed_purchase_orders(
                user
            )
            .get(
                pk=purchase_order.pk
            )
        )

        return Response(
            PurchaseOrderDetailSerializer(
                purchase_order
            ).data,
            status=status.HTTP_201_CREATED,
        )


# ============================================================
# PURCHASE ORDER DETAIL
# ============================================================

class PurchaseOrderDetailView(APIView):

    permission_classes = [
        IsAuthenticated
    ]

    def get_object(
        self,
        request,
        pk
    ):

        try:

            return get_allowed_purchase_orders(
                request.user
            ).get(
                pk=pk
            )

        except PurchaseOrder.DoesNotExist:

            return None

    def get(
        self,
        request,
        pk
    ):

        purchase_order = self.get_object(
            request,
            pk
        )

        if purchase_order is None:

            return Response(
                {
                    "detail":
                    "Purchase Order not found."
                },
                status=status.HTTP_404_NOT_FOUND,
            )

        return Response(
            PurchaseOrderDetailSerializer(
                purchase_order
            ).data
        )

    def delete(
        self,
        request,
        pk
    ):

        purchase_order = self.get_object(
            request,
            pk
        )

        if purchase_order is None:

            return Response(
                {
                    "detail":
                    "Purchase Order not found."
                },
                status=status.HTTP_404_NOT_FOUND,
            )

        if purchase_order.status != "draft":

            return Response(
                {
                    "detail":
                    "Only Draft Purchase Orders can be deleted."
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        purchase_order.delete()

        return Response(
            {
                "detail":
                "Purchase Order deleted successfully."
            },
            status=status.HTTP_204_NO_CONTENT,
        )


# ============================================================
# PURCHASE ORDER ORDER / CONFIRM
# ============================================================

class PurchaseOrderOrderView(APIView):

    permission_classes = [
        IsAuthenticated
    ]

    def post(
        self,
        request,
        pk
    ):

        try:

            purchase_order = (
                get_allowed_purchase_orders(
                    request.user
                ).get(
                    pk=pk
                )
            )

        except PurchaseOrder.DoesNotExist:

            return Response(
                {
                    "detail":
                    "Purchase Order not found."
                },
                status=status.HTTP_404_NOT_FOUND,
            )

        if purchase_order.status != "draft":

            return Response(
                {
                    "detail":
                    "Only Draft Purchase Orders can be ordered."
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        if not purchase_order.items.exists():

            return Response(
                {
                    "detail":
                    "Purchase Order must contain at least one product."
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        purchase_order.status = "ordered"

        purchase_order.save(
            update_fields=[
                "status",
                "updated_at",
            ]
        )

        return Response(
            PurchaseOrderDetailSerializer(
                purchase_order
            ).data
        )


# ============================================================
# PURCHASE ORDER CANCEL
# ============================================================

class PurchaseOrderCancelView(APIView):

    permission_classes = [
        IsAuthenticated
    ]

    def post(
        self,
        request,
        pk
    ):

        try:

            purchase_order = (
                get_allowed_purchase_orders(
                    request.user
                ).get(
                    pk=pk
                )
            )

        except PurchaseOrder.DoesNotExist:

            return Response(
                {
                    "detail":
                    "Purchase Order not found."
                },
                status=status.HTTP_404_NOT_FOUND,
            )

        if purchase_order.status in (
            "fully_received",
            "cancelled",
        ):

            return Response(
                {
                    "detail":
                    "This Purchase Order cannot be cancelled."
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        purchase_order.status = "cancelled"

        purchase_order.save(
            update_fields=[
                "status",
                "updated_at",
            ]
        )

        return Response(
            PurchaseOrderDetailSerializer(
                purchase_order
            ).data
        )


# ============================================================
# PURCHASE ORDER UPDATE
# ============================================================

class PurchaseOrderUpdateView(APIView):

    permission_classes = [
        IsAuthenticated
    ]

    def put(
        self,
        request,
        pk
    ):

        return self.update(
            request,
            pk,
            partial=False
        )

    def patch(
        self,
        request,
        pk
    ):

        return self.update(
            request,
            pk,
            partial=True
        )

    def update(
        self,
        request,
        pk,
        partial=False
    ):

        try:

            purchase_order = (
                get_allowed_purchase_orders(
                    request.user
                ).get(
                    pk=pk
                )
            )

        except PurchaseOrder.DoesNotExist:

            return Response(
                {
                    "detail":
                    "Purchase Order not found."
                },
                status=status.HTTP_404_NOT_FOUND,
            )

        if purchase_order.status != "draft":

            return Response(
                {
                    "detail":
                    "Only Draft Purchase Orders can be edited."
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        serializer = PurchaseOrderWriteSerializer(
            data=request.data,
            partial=partial
        )

        if not serializer.is_valid():

            return Response(
                serializer.errors,
                status=status.HTTP_400_BAD_REQUEST,
            )

        validated = serializer.validated_data

        branch_id = validated.get(
            "branch_id"
        )

        if branch_id:

            branch = resolve_product_branch(
                request.user,
                branch_id
            )

        else:

            branch = purchase_order.branch

        if branch is None:

            return Response(
                {
                    "detail":
                    "Invalid branch."
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        supplier = validated.get(
            "supplier",
            purchase_order.supplier
        )

        if supplier.status != "active":

            return Response(
                {
                    "detail":
                    "Inactive supplier cannot be used."
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        items = validated.get(
            "items"
        )

        if items is not None:

            if not items:

                return Response(
                    {
                        "detail":
                        "At least one product is required."
                    },
                    status=status.HTTP_400_BAD_REQUEST,
                )

            for item in items:

                if (
                    item["product"].branch_id
                    != branch.id
                ):

                    return Response(
                        {
                            "detail":
                            (
                                f"Product "
                                f"{item['product'].barcode} "
                                f"does not belong to selected branch."
                            )
                        },
                        status=status.HTTP_400_BAD_REQUEST,
                    )

        with transaction.atomic():

            purchase_order.branch = branch

            purchase_order.supplier = supplier

            if "po_date" in validated:

                purchase_order.po_date = validated[
                    "po_date"
                ]

            if (
                "expected_delivery_date"
                in validated
            ):

                purchase_order.expected_delivery_date = (
                    validated[
                        "expected_delivery_date"
                    ]
                )

            purchase_order.save()

            if items is not None:

                purchase_order.items.all().delete()

                for item in items:

                    PurchaseOrderItem.objects.create(
                        purchase_order=purchase_order,
                        product=item["product"],
                        quantity=item["quantity"],
                        purchase_price=item[
                            "purchase_price"
                        ],
                    )

            if "containers" in validated:

                purchase_order.containers.all().delete()

                for container in validated[
                    "containers"
                ]:

                    PurchaseOrderContainer.objects.create(
                        purchase_order=purchase_order,
                        container_number=container[
                            "container_number"
                        ],
                        container_type=container.get(
                            "container_type",
                            ""
                        ),
                        quantity=container[
                            "quantity"
                        ],
                        remarks=container.get(
                            "remarks",
                            ""
                        ),
                    )

        purchase_order.refresh_from_db()

        return Response(
            PurchaseOrderDetailSerializer(
                purchase_order
            ).data
        )


# ============================================================
# RECEIVING HELPER
# ============================================================

def update_purchase_order_status(
    purchase_order
):

    items = list(
        purchase_order.items.all()
    )

    if not items:

        purchase_order.status = "draft"

        purchase_order.save(
            update_fields=[
                "status",
                "updated_at",
            ]
        )

        return

    total_ordered = sum(
        item.quantity
        for item in items
    )

    total_received = sum(
        item.received_quantity
        for item in items
    )

    if total_received == 0:

        purchase_order.status = "ordered"

    elif total_received < total_ordered:

        purchase_order.status = (
            "partially_received"
        )

    else:

        purchase_order.status = (
            "fully_received"
        )

    purchase_order.save(
        update_fields=[
            "status",
            "updated_at",
        ]
    )


# ============================================================
# PURCHASE RECEIVING
# ============================================================

class PurchaseOrderReceiveView(APIView):

    permission_classes = [
        IsAuthenticated
    ]

    def post(
        self,
        request,
        pk
    ):

        try:

            purchase_order = (
                get_allowed_purchase_orders(
                    request.user
                )
                .prefetch_related(
                    "items__product"
                )
                .get(
                    pk=pk
                )
            )

        except PurchaseOrder.DoesNotExist:

            return Response(
                {
                    "detail":
                    "Purchase Order not found."
                },
                status=status.HTTP_404_NOT_FOUND,
            )

        if purchase_order.status in (
            "draft",
            "cancelled",
            "fully_received",
        ):

            return Response(
                {
                    "detail":
                    "This Purchase Order cannot receive products in its current status."
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        receipt_date = request.data.get(
            "receipt_date"
        )

        if not receipt_date:

            return Response(
                {
                    "detail":
                    "receipt_date is required."
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        remarks = request.data.get(
            "remarks",
            ""
        )

        items_data = request.data.get(
            "items"
        )

        if not isinstance(
            items_data,
            list
        ) or not items_data:

            return Response(
                {
                    "detail":
                    "At least one receiving item is required."
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        po_items = {
            item.id: item
            for item in purchase_order.items.all()
        }

        receipt_items = []

        try:

            with transaction.atomic():

                receipt = PurchaseReceipt.objects.create(
                    purchase_order=purchase_order,
                    receipt_date=receipt_date,
                    remarks=remarks,
                    received_by=request.user,
                )

                for data in items_data:

                    item_id = data.get(
                        "purchase_order_item"
                    )

                    quantity = data.get(
                        "quantity"
                    )

                    if not item_id:

                        raise ValueError(
                            "purchase_order_item is required."
                        )

                    if quantity is None:

                        raise ValueError(
                            "Receiving quantity is required."
                        )

                    try:

                        quantity = int(
                            quantity
                        )

                    except (
                        ValueError,
                        TypeError,
                    ):

                        raise ValueError(
                            "Receiving quantity must be an integer."
                        )

                    if quantity <= 0:

                        raise ValueError(
                            "Receiving quantity must be greater than zero."
                        )

                    try:

                        po_item = po_items[
                            int(item_id)
                        ]

                    except (
                        KeyError,
                        ValueError,
                        TypeError,
                    ):

                        raise ValueError(
                            "Invalid Purchase Order Item."
                        )

                    current_received = (
                        po_item.received_quantity
                    )

                    pending = (
                        po_item.quantity
                        - current_received
                    )

                    if quantity > pending:

                        raise ValueError(
                            (
                                f"Cannot receive {quantity} "
                                f"for {po_item.product.barcode}. "
                                f"Pending quantity is {pending}."
                            )
                        )

                    PurchaseReceiptItem.objects.create(
                        receipt=receipt,
                        purchase_order_item=po_item,
                        quantity=quantity,
                    )

                    po_item.received_quantity = (
                        current_received
                        + quantity
                    )

                    po_item.save(
                        update_fields=[
                            "received_quantity",
                            "updated_at",
                        ]
                    )

                    receipt_items.append(
                        po_item
                    )

                update_purchase_order_status(
                    purchase_order
                )

        except Exception as exc:

            return Response(
                {
                    "detail":
                    str(exc)
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        receipt.refresh_from_db()

        return Response(
            PurchaseReceiptSerializer(
                receipt
            ).data,
            status=status.HTTP_201_CREATED,
        )


# ============================================================
# RECEIVING HISTORY
# ============================================================

class PurchaseOrderReceivingHistoryView(
    APIView
):

    permission_classes = [
        IsAuthenticated
    ]

    def get(
        self,
        request,
        pk
    ):

        try:

            purchase_order = (
                get_allowed_purchase_orders(
                    request.user
                )
                .get(
                    pk=pk
                )
            )

        except PurchaseOrder.DoesNotExist:

            return Response(
                {
                    "detail":
                    "Purchase Order not found."
                },
                status=status.HTTP_404_NOT_FOUND,
            )

        receipts = (
            PurchaseReceipt.objects
            .filter(
                purchase_order=purchase_order
            )
            .prefetch_related(
                "items__purchase_order_item__product"
            )
            .order_by(
                "-receipt_date",
                "-id",
            )
        )

        return Response(
            PurchaseReceiptSerializer(
                receipts,
                many=True
            ).data
        )