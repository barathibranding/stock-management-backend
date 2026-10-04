from django.contrib import admin
from django.contrib.auth import get_user_model
from django.contrib.auth.admin import UserAdmin
from django.http import HttpResponse
from django.urls import path
from django.shortcuts import render, redirect
from django.contrib import messages

from openpyxl import load_workbook

from .models import (
    Branch,
    UserProfile,
    Product,
    Supplier,
    PurchaseOrder,
    PurchaseOrderItem,
    PurchaseOrderContainer,
    PurchaseReceipt,
    PurchaseReceiptItem,
)


User = get_user_model()


# ============================================================
# BRANCH ADMIN
# ============================================================

@admin.register(Branch)
class BranchAdmin(admin.ModelAdmin):
    list_display = (
        "id",
        "branch_name",
        "created_at",
        "updated_at",
    )

    search_fields = (
        "branch_name",
    )

    list_per_page = 25


# ============================================================
# USER PROFILE INLINE
# ============================================================

class UserProfileInline(admin.StackedInline):
    model = UserProfile

    extra = 1
    max_num = 1
    can_delete = True

    autocomplete_fields = (
        "branch",
    )

    def get_formset(
        self,
        request,
        obj=None,
        **kwargs
    ):
        formset = super().get_formset(
            request,
            obj,
            **kwargs
        )

        original_init = formset.form.__init__

        def custom_init(
            form,
            *args,
            **form_kwargs
        ):
            original_init(
                form,
                *args,
                **form_kwargs
            )

            user_type = form.data.get(
                form.add_prefix(
                    "user_type"
                ),
                form.initial.get(
                    "user_type",
                    "staff"
                ),
            )

            if user_type == "admin":
                form.fields[
                    "branch"
                ].required = False

        formset.form.__init__ = custom_init

        return formset


# ============================================================
# USER ADMIN
# ============================================================

admin.site.unregister(User)


@admin.register(User)
class CustomUserAdmin(UserAdmin):

    inlines = [
        UserProfileInline
    ]

    list_display = (
        "username",
        "first_name",
        "last_name",
        "email",
        "is_active",
        "is_staff",
    )

    list_filter = (
        "is_active",
        "is_staff",
        "is_superuser",
        "groups",
    )

    search_fields = (
        "username",
        "first_name",
        "last_name",
        "email",
    )


# ============================================================
# PRODUCT ADMIN
# ============================================================

@admin.action(
    description="Export selected products to Excel"
)
def export_selected_products(
    modeladmin,
    request,
    queryset
):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from io import BytesIO

    workbook = Workbook()

    worksheet = workbook.active
    worksheet.title = "Inventory"

    headers = [
        "S.R",
        "BarCode",
        "Short Description",
        "Location",
        "Cost Price",
        "Selling Price",
        "Stk On Hand",
        "Physical Qty",
        "Variance",
        "Branch",
    ]

    worksheet.append(headers)

    for cell in worksheet[1]:
        cell.font = Font(
            bold=True
        )

    for index, product in enumerate(
        queryset.select_related(
            "branch"
        ),
        start=1
    ):
        worksheet.append([
            index,
            product.barcode,
            product.short_description,
            product.location,
            float(product.cost_price),
            float(product.selling_price),
            product.stk_on_hand,
            product.physical_qty,
            product.variance,
            product.branch.branch_name,
        ])

    for column, width in {
        "A": 10,
        "B": 25,
        "C": 45,
        "D": 20,
        "E": 15,
        "F": 15,
        "G": 15,
        "H": 15,
        "I": 15,
        "J": 25,
    }.items():
        worksheet.column_dimensions[
            column
        ].width = width

    output = BytesIO()

    workbook.save(output)

    output.seek(0)

    response = HttpResponse(
        output.getvalue(),
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
    )

    response[
        "Content-Disposition"
    ] = (
        'attachment; '
        'filename="selected_inventory.xlsx"'
    )

    return response


@admin.register(Product)
class ProductAdmin(admin.ModelAdmin):

    list_display = (
        "barcode",
        "short_description",
        "branch",
        "location",
        "cost_price",
        "selling_price",
        "stk_on_hand",
        "physical_qty",
        "variance_display",
        "updated_at",
    )

    list_filter = (
        "branch",
        "location",
    )

    search_fields = (
        "barcode",
        "short_description",
        "location",
    )

    ordering = (
        "barcode",
    )

    list_per_page = 50

    autocomplete_fields = (
        "branch",
    )

    readonly_fields = (
        "variance_display",
        "created_at",
        "updated_at",
    )

    fieldsets = (
        (
            "Product Information",
            {
                "fields": (
                    "branch",
                    "barcode",
                    "short_description",
                    "location",
                )
            }
        ),
        (
            "Pricing",
            {
                "fields": (
                    "cost_price",
                    "selling_price",
                )
            }
        ),
        (
            "Stock",
            {
                "fields": (
                    "stk_on_hand",
                    "physical_qty",
                    "variance_display",
                )
            }
        ),
        (
            "System",
            {
                "fields": (
                    "created_at",
                    "updated_at",
                )
            }
        ),
    )

    actions = [
        export_selected_products,
    ]

    def variance_display(self, obj):
        return obj.variance

    variance_display.short_description = (
        "Variance"
    )

    def get_urls(self):
        urls = super().get_urls()

        custom_urls = [
            path(
                "import-excel/",
                self.admin_site.admin_view(
                    self.import_excel_view
                ),
                name="inventory_product_import_excel",
            ),
        ]

        return custom_urls + urls

    def import_excel_view(
        self,
        request
    ):
        if request.method == "POST":

            excel_file = request.FILES.get(
                "excel_file"
            )

            if not excel_file:
                messages.error(
                    request,
                    "Please select an Excel file."
                )

                return redirect(
                    request.path
                )

            branch_id = request.POST.get(
                "branch"
            )

            if not branch_id:
                messages.error(
                    request,
                    "Please select a branch."
                )

                return redirect(
                    request.path
                )

            try:
                branch = Branch.objects.get(
                    pk=branch_id
                )

            except Branch.DoesNotExist:
                messages.error(
                    request,
                    "Selected branch does not exist."
                )

                return redirect(
                    request.path
                )

            if not excel_file.name.lower().endswith(
                ".xlsx"
            ):
                messages.error(
                    request,
                    "Only .xlsx files are supported."
                )

                return redirect(
                    request.path
                )

            try:
                workbook = load_workbook(
                    excel_file,
                    read_only=True,
                    data_only=True,
                )

                if "Inventory" in workbook.sheetnames:
                    worksheet = workbook[
                        "Inventory"
                    ]
                else:
                    worksheet = workbook[
                        workbook.sheetnames[0]
                    ]

                rows = worksheet.iter_rows(
                    values_only=True
                )

                headers = next(rows)

                header_map = {
                    str(value).strip():
                    index
                    for index, value in enumerate(
                        headers
                    )
                    if value is not None
                }

                required = [
                    "BarCode",
                    "Short Description",
                    "Location",
                    "Cost Price",
                    "Selling Price",
                    "Stk On Hand",
                    "Physical Qty",
                ]

                missing = [
                    column
                    for column in required
                    if column not in header_map
                ]

                if missing:
                    messages.error(
                        request,
                        "Missing columns: "
                        + ", ".join(missing)
                    )

                    return redirect(
                        request.path
                    )

                existing = {
                    product.barcode:
                    product
                    for product in Product.objects.filter(
                        branch=branch
                    )
                }

                create_list = []
                update_list = []

                seen = set()

                errors = []

                from decimal import Decimal

                for row_number, row in enumerate(
                    rows,
                    start=2
                ):
                    try:
                        barcode = str(
                            row[
                                header_map["BarCode"]
                            ] or ""
                        ).strip()

                        if not barcode:
                            raise ValueError(
                                "Barcode is empty."
                            )

                        if barcode in seen:
                            raise ValueError(
                                "Duplicate barcode in file."
                            )

                        seen.add(barcode)

                        description = str(
                            row[
                                header_map[
                                    "Short Description"
                                ]
                            ] or ""
                        ).strip()

                        location = str(
                            row[
                                header_map["Location"]
                            ] or ""
                        ).strip()

                        cost_price = Decimal(
                            str(
                                row[
                                    header_map[
                                        "Cost Price"
                                    ]
                                ] or 0
                            )
                        )

                        selling_price = Decimal(
                            str(
                                row[
                                    header_map[
                                        "Selling Price"
                                    ]
                                ] or 0
                            )
                        )

                        stk_on_hand = int(
                            float(
                                row[
                                    header_map[
                                        "Stk On Hand"
                                    ]
                                ] or 0
                            )
                        )

                        physical_qty = int(
                            float(
                                row[
                                    header_map[
                                        "Physical Qty"
                                    ]
                                ] or 0
                            )
                        )

                        product = existing.get(
                            barcode
                        )

                        if product:
                            product.short_description = (
                                description
                            )
                            product.location = (
                                location
                            )
                            product.cost_price = (
                                cost_price
                            )
                            product.selling_price = (
                                selling_price
                            )
                            product.stk_on_hand = (
                                stk_on_hand
                            )
                            product.physical_qty = (
                                physical_qty
                            )

                            update_list.append(
                                product
                            )

                        else:
                            create_list.append(
                                Product(
                                    branch=branch,
                                    barcode=barcode,
                                    short_description=(
                                        description
                                    ),
                                    location=location,
                                    cost_price=(
                                        cost_price
                                    ),
                                    selling_price=(
                                        selling_price
                                    ),
                                    stk_on_hand=(
                                        stk_on_hand
                                    ),
                                    physical_qty=(
                                        physical_qty
                                    ),
                                )
                            )

                    except Exception as exc:
                        errors.append(
                            f"Row {row_number}: {exc}"
                        )

                if errors:
                    for error in errors[:20]:
                        messages.error(
                            request,
                            error
                        )

                    return redirect(
                        request.path
                    )

                from django.db import transaction

                with transaction.atomic():

                    if create_list:
                        Product.objects.bulk_create(
                            create_list,
                            batch_size=1000
                        )

                    if update_list:
                        Product.objects.bulk_update(
                            update_list,
                            fields=[
                                "short_description",
                                "location",
                                "cost_price",
                                "selling_price",
                                "stk_on_hand",
                                "physical_qty",
                            ],
                            batch_size=1000
                        )

                messages.success(
                    request,
                    (
                        f"Import completed. "
                        f"Created: {len(create_list)}, "
                        f"Updated: {len(update_list)}"
                    )
                )

                return redirect(
                    "admin:inventory_product_changelist"
                )

            except Exception as exc:
                messages.error(
                    request,
                    f"Import failed: {exc}"
                )

                return redirect(
                    request.path
                )

        branches = Branch.objects.all()

        context = {
            **self.admin_site.each_context(
                request
            ),
            "title":
            "Import Products from Excel",
            "branches":
            branches,
        }

        return render(
            request,
            "admin/inventory/product/import_excel.html",
            context
        )
        
# ============================================================
# SUPPLIER ADMIN
# ============================================================

@admin.register(Supplier)
class SupplierAdmin(admin.ModelAdmin):

    list_display = (
        "id",
        "supplier_name",
        "supplier_location",
        "status",
        "created_at",
        "updated_at",
    )

    list_filter = (
        "status",
    )

    search_fields = (
        "supplier_name",
        "supplier_location",
    )

    ordering = (
        "supplier_name",
    )

    list_per_page = 50


# ============================================================
# PURCHASE ORDER ITEM INLINE
# ============================================================

class PurchaseOrderItemInline(
    admin.TabularInline
):

    model = PurchaseOrderItem

    extra = 0

    autocomplete_fields = (
        "product",
    )

    readonly_fields = (
        "received_quantity",
    )


# ============================================================
# PURCHASE ORDER CONTAINER INLINE
# ============================================================

class PurchaseOrderContainerInline(
    admin.TabularInline
):

    model = PurchaseOrderContainer

    extra = 0


# ============================================================
# PURCHASE ORDER ADMIN
# ============================================================

@admin.register(PurchaseOrder)
class PurchaseOrderAdmin(admin.ModelAdmin):

    list_display = (
        "po_number",
        "branch",
        "supplier",
        "po_date",
        "expected_delivery_date",
        "status",
        "total_products_display",
        "total_quantity_display",
        "total_amount_display",
        "created_at",
    )

    list_filter = (
        "status",
        "branch",
        "supplier",
        "po_date",
        "expected_delivery_date",
    )

    search_fields = (
        "po_number",
        "supplier__supplier_name",
    )

    ordering = (
        "-created_at",
    )

    list_per_page = 50

    autocomplete_fields = (
        "supplier",
        "branch",
        "created_by",
    )

    readonly_fields = (
        "po_number",
        "created_by",
        "created_at",
        "updated_at",
        "total_products_display",
        "total_quantity_display",
        "total_received_quantity_display",
        "total_boxes_display",
        "total_amount_display",
    )

    fieldsets = (
        (
            "Purchase Order",
            {
                "fields": (
                    "po_number",
                    "branch",
                    "supplier",
                    "po_date",
                    "expected_delivery_date",
                    "status",
                )
            }
        ),

        (
            "Summary",
            {
                "fields": (
                    "total_products_display",
                    "total_quantity_display",
                    "total_received_quantity_display",
                    "total_boxes_display",
                    "total_amount_display",
                )
            }
        ),

        (
            "System",
            {
                "fields": (
                    "created_by",
                    "created_at",
                    "updated_at",
                )
            }
        ),
    )

    inlines = [
        PurchaseOrderItemInline,
        PurchaseOrderContainerInline,
    ]

    def total_products_display(
        self,
        obj
    ):
        return obj.total_products

    total_products_display.short_description = (
        "Total Products"
    )

    def total_quantity_display(
        self,
        obj
    ):
        return obj.total_quantity

    total_quantity_display.short_description = (
        "Total Quantity"
    )

    def total_received_quantity_display(
        self,
        obj
    ):
        return obj.total_received_quantity

    total_received_quantity_display.short_description = (
        "Received Quantity"
    )

    def total_boxes_display(
        self,
        obj
    ):
        return obj.total_boxes

    total_boxes_display.short_description = (
        "Total Boxes"
    )

    def total_amount_display(
        self,
        obj
    ):
        return obj.total_amount

    total_amount_display.short_description = (
        "Total Amount"
    )


# ============================================================
# PURCHASE ORDER ITEM ADMIN
# ============================================================

@admin.register(PurchaseOrderItem)
class PurchaseOrderItemAdmin(
    admin.ModelAdmin
):

    list_display = (
        "purchase_order",
        "product",
        "quantity",
        "purchase_price",
        "total_amount_display",
        "received_quantity",
        "pending_quantity_display",
    )

    search_fields = (
        "purchase_order__po_number",
        "product__barcode",
        "product__short_description",
    )

    autocomplete_fields = (
        "purchase_order",
        "product",
    )

    readonly_fields = (
        "total_amount_display",
        "pending_quantity_display",
    )

    def total_amount_display(
        self,
        obj
    ):
        return obj.total_amount

    total_amount_display.short_description = (
        "Total Amount"
    )

    def pending_quantity_display(
        self,
        obj
    ):
        return obj.pending_quantity

    pending_quantity_display.short_description = (
        "Pending Quantity"
    )


# ============================================================
# CONTAINER ADMIN
# ============================================================

@admin.register(PurchaseOrderContainer)
class PurchaseOrderContainerAdmin(
    admin.ModelAdmin
):

    list_display = (
        "purchase_order",
        "container_number",
        "container_type",
        "quantity",
        "remarks",
        "created_at",
    )

    search_fields = (
        "purchase_order__po_number",
        "container_number",
        "container_type",
    )

    list_filter = (
        "container_type",
    )

    autocomplete_fields = (
        "purchase_order",
    )


# ============================================================
# RECEIPT ITEM INLINE
# ============================================================

class PurchaseReceiptItemInline(
    admin.TabularInline
):

    model = PurchaseReceiptItem

    extra = 0

    autocomplete_fields = (
        "purchase_order_item",
    )


# ============================================================
# RECEIPT ADMIN
# ============================================================

@admin.register(PurchaseReceipt)
class PurchaseReceiptAdmin(admin.ModelAdmin):

    list_display = (
        "id",
        "purchase_order",
        "receipt_date",
        "received_by",
        "total_quantity_display",
        "created_at",
    )

    list_filter = (
        "receipt_date",
    )

    search_fields = (
        "purchase_order__po_number",
        "received_by__username",
        "received_by__first_name",
        "received_by__last_name",
    )

    autocomplete_fields = (
        "purchase_order",
        "received_by",
    )

    readonly_fields = (
        "created_at",
        "total_quantity_display",
    )

    inlines = [
        PurchaseReceiptItemInline,
    ]

    def total_quantity_display(
        self,
        obj
    ):
        return obj.total_quantity

    total_quantity_display.short_description = (
        "Total Received Quantity"
    )


# ============================================================
# RECEIPT ITEM ADMIN
# ============================================================

@admin.register(PurchaseReceiptItem)
class PurchaseReceiptItemAdmin(
    admin.ModelAdmin
):

    list_display = (
        "receipt",
        "purchase_order_item",
        "quantity",
        "created_at",
    )

    search_fields = (
        "receipt__purchase_order__po_number",
        "purchase_order_item__product__barcode",
    )

    autocomplete_fields = (
        "receipt",
        "purchase_order_item",
    )        